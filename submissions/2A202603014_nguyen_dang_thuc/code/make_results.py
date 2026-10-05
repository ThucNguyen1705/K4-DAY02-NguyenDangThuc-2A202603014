"""make_results.py - Bước 5: gom log của các bước thành results.xlsx (7 sheet), biểu đồ tổng hợp
và report_tables.md (bảng markdown để chép vào báo cáo). Mọi số đều đọc từ file log/predictions.

    python make_results.py --out <thư mục output>
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from dataset import CLASS_NAMES
from train import compute_metrics

HARD = [0, 7]


def _fmt_ms(mean, std, d=4):
    return f"{mean:.{d}f} ± {std:.{d}f}"


def backbones_sheet(root: Path) -> pd.DataFrame:
    b = pd.read_csv(root / "step1_backbones.csv")
    lat = pd.read_csv(root / "step3_latency.csv")
    l1 = lat[(lat["batch"] == 1) & (lat["method"].isin(b["exp_id"]))].set_index("method")
    l32 = lat[(lat["batch"] == 32) & (lat["method"].isin(b["exp_id"]))].set_index("method")
    f1 = b["val_f1_per_class"].map(_parse_list)
    return pd.DataFrame({
        "exp_id": b["exp_id"], "backbone": b["backbone"], "tag trọng số": b["weight_tag"],
        "#tham số (M)": b["params_m"], "GMAC": b["gmacs"], "độ phân giải": b["img_size"],
        "epoch": b["epochs"], "seed": b["seed"], "best epoch": b["best_epoch"],
        "macro-F1 val": b["val_macro_f1"], "top-1 val": b["val_top1"],
        "F1 Chinee apple val": [r[0] for r in f1], "F1 Snake weed val": [r[7] for r in f1],
        "thời gian train/epoch (s)": b["sec_per_epoch"],
        "độ trễ batch-1 p50 (ms)": b["exp_id"].map(l1["p50"]),
        "độ trễ batch-1 p95 (ms)": b["exp_id"].map(l1["p95"]),
        "thông lượng batch-32 (ảnh/s)": b["exp_id"].map(l32["images_per_s"]),
        "ghi chú": "công thức nền T00, 1 seed; độ trễ FP32 đo ở Bước 3 (warmup 10, synchronize, 60 lần)",
    })


def _parse_list(v):
    return json.loads(v) if isinstance(v, str) else v


def training_sheet(root: Path) -> pd.DataFrame:
    t = pd.read_csv(root / "step2_training.csv")
    f1 = t["val_f1_per_class"].map(_parse_list)
    note = np.where(t["distinguishable"], "Δ vượt std nhiễu", "không phân biệt được với T00 (|Δ| ≤ std)")
    note = np.where(t["exp_id"] == "T00", "công thức nền; std qua 3 seed = " + t["noise_std"].map("{:.4f}".format), note)
    return pd.DataFrame({
        "exp_id": t["exp_id"], "backbone": t["backbone"], "trục (A–G)": t["axis"],
        "khác T00 ở điểm nào": t["change"], "seed": t["seed"], "best epoch": t["best_epoch"],
        "macro-F1 val": t["val_macro_f1"], "top-1 val": t["val_top1"], "Δ macro-F1 so với T00": t["delta_vs_T00"],
        "std nhiễu T00 (3 seed)": t["noise_std"],
        "F1 Chinee apple val": [r[0] for r in f1], "F1 Snake weed val": [r[7] for r in f1],
        "F1 lớp thấp nhất val": [min(r) for r in f1],
        "F1 Negatives val": [r[8] for r in f1],
        "thời gian train/epoch (s)": t["sec_per_epoch"], "ghi chú": note,
    })


def inference_sheet(root: Path) -> pd.DataFrame:
    i = pd.read_csv(root / "step3_inference.csv")
    cols = {"exp_id": "exp_id", "method": "phương pháp", "model": "mô hình/checkpoint", "K": "K (view hoặc mô hình)",
            "val_macro_f1": "macro-F1 val", "val_top1": "top-1 val", "val_ece": "ECE val", "val_nll": "NLL val",
            "f1_chinee": "F1 Chinee apple val", "f1_snake": "F1 Snake weed val",
            "lat_b1_p50_ms": "p50 batch-1 (ms)", "lat_b1_p95_ms": "p95 batch-1 (ms)", "lat_b1_p99_ms": "p99 batch-1 (ms)",
            "throughput_b32_img_s": "thông lượng batch-32 (ảnh/s)", "rel_cost_vs_I00": "chi phí tương đối so với I00",
            "ece_before": "ECE val trước TS", "ece_after_crossfit": "ECE val sau TS (khớp chéo 2 nửa)", "T": "T", "note": "ghi chú"}
    return i[[c for c in cols if c in i.columns]].rename(columns=cols)


def latency_sheet(root: Path) -> pd.DataFrame:
    l = pd.read_csv(root / "step3_latency.csv")
    out = pd.DataFrame({
        "cấu hình": l["config"], "phương pháp": l["method"], "GPU": l["gpu"], "torch": l["torch"],
        "dtype": l["dtype"], "batch": l["batch"], "độ phân giải": l["img_size"],
        "gộp BN": l["fused_bn"].map({True: "có", False: "không"}),
        "tính tiền xử lý": l["preprocess"].map({True: "có", False: "không"}),
        "p50 (ms)": l["p50"], "p95 (ms)": l["p95"], "p99 (ms)": l["p99"], "số lần đo": l["n"],
        "ảnh/s": l["images_per_s"],
    })
    return out


def final_sheet(root: Path, dec: dict) -> pd.DataFrame:
    fv = pd.read_csv(root / "step4_final_val.csv")
    rec, meth = dec["step2_recipe"], dec["step3_method"]
    desc = {"F01": f"{rec['backbone']} + công thức {rec['exp_id']} {rec['cfg'] or '(= T00)'} + suy luận {meth['method']} + TS",
            "T00": f"{rec['backbone']} + công thức nền T00 + suy luận I00 (1 view)"}
    rows = []
    for tag in ("F01", "T00"):
        ps = pd.read_csv(root / "eval_out" / f"{tag}_per_seed.csv")
        summ = json.loads((root / "eval_out" / f"{tag}_summary.json").read_text())
        v = fv[fv["exp_id"] == tag].set_index("seed")
        for _, r in ps.iterrows():
            rows.append({"exp_id": tag, "cấu hình": desc[tag], "seed": int(r["seed"]),
                         "macro-F1 val": v.loc[int(r["seed"]), "val_macro_f1"],
                         "macro-F1 test": r["macro_f1"], "top-1 test": r["top1"],
                         "balanced acc test": r["balanced_acc"], "ECE test": r["ece"], "file": r["file"]})
        vals = v["val_macro_f1"]
        rows.append({"exp_id": tag, "cấu hình": desc[tag] + " — mean ± std", "seed": f"{len(ps)} seed",
                     "macro-F1 val": _fmt_ms(vals.mean(), vals.std(ddof=1)),
                     "macro-F1 test": _fmt_ms(summ["macro_f1"]["mean"], summ["macro_f1"]["std"]),
                     "top-1 test": _fmt_ms(summ["top1"]["mean"], summ["top1"]["std"]),
                     "balanced acc test": _fmt_ms(summ["balanced_acc"]["mean"], summ["balanced_acc"]["std"]),
                     "ECE test": _fmt_ms(summ["ece"]["mean"], summ["ece"]["std"]), "file": "eval.py score"})
    return pd.DataFrame(rows)


def perclass_sheet(root: Path) -> pd.DataFrame:
    rows = []
    fv = pd.read_csv(root / "step4_final_val.csv")
    best_seed = int(fv[fv["exp_id"] == "F01"].sort_values("val_macro_f1", ascending=False)["seed"].iloc[0])
    for tag, label in (("F01", "chung kết F01 (mean ± std 3 seed)"), ("T00", "mốc T00+I00 (mean ± std 3 seed)")):
        pc = pd.read_csv(root / "eval_out" / f"{tag}_per_class.csv")
        for _, r in pc.iterrows():
            rows.append({"cấu hình": label, "lớp": r["class"], "số ảnh test": int(r["support"]),
                         "precision": _fmt_ms(r["precision_mean"], r["precision_std"]),
                         "recall": _fmt_ms(r["recall_mean"], r["recall_std"]),
                         "F1": _fmt_ms(r["f1_mean"], r["f1_std"])})
    p = pd.read_csv(root / "predictions" / f"F01_seed{best_seed}_test.csv")
    probs = p[[f"p{i}" for i in range(9)]].to_numpy()
    m = compute_metrics(p["y_true"].to_numpy(), p["y_pred"].to_numpy(), probs)
    for i, name in enumerate(CLASS_NAMES):
        rows.append({"cấu hình": f"F01 seed {best_seed} (seed có macro-F1 VAL cao nhất)", "lớp": name,
                     "số ảnh test": int(m["support"][i]), "precision": f"{m['precision'][i]:.4f}",
                     "recall": f"{m['recall'][i]:.4f}", "F1": f"{m['f1'][i]:.4f}"})
    return pd.DataFrame(rows)


def summary_sheet(bk, tr, inf, fin) -> pd.DataFrame:
    rows = []
    for _, r in bk.iterrows():
        rows.append({"exp_id": r["exp_id"], "loại": "backbone", "mô tả": r["backbone"], "macro-F1 val": r["macro-F1 val"],
                     "top-1 val": r["top-1 val"], "#tham số (M)": r["#tham số (M)"], "GMAC": r["GMAC"],
                     "p50 batch-1 (ms)": r["độ trễ batch-1 p50 (ms)"], "train s/epoch": r["thời gian train/epoch (s)"]})
    for _, r in tr[tr["seed"] == 0].iterrows():
        rows.append({"exp_id": r["exp_id"], "loại": f"công thức (trục {r['trục (A–G)']})", "mô tả": r["khác T00 ở điểm nào"],
                     "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                     "train s/epoch": r["thời gian train/epoch (s)"]})
    for _, r in inf.dropna(subset=["macro-F1 val"]).iterrows():
        if r["exp_id"] in ("I00", "I07", "I08_unfused"):
            continue  # trùng số với mô hình gốc
        rows.append({"exp_id": r["exp_id"], "loại": "suy luận", "mô tả": r["phương pháp"], "macro-F1 val": r["macro-F1 val"],
                     "top-1 val": r["top-1 val"], "p50 batch-1 (ms)": r.get("p50 batch-1 (ms)"),
                     "chi phí so với I00": r.get("chi phí tương đối so với I00")})
    df = pd.DataFrame(rows).sort_values("macro-F1 val", ascending=False).head(10).reset_index(drop=True)
    df.insert(0, "hạng", range(1, len(df) + 1))
    f = fin[fin["seed"].astype(str).str.contains("seed")]
    for _, r in f.iterrows():
        df = pd.concat([df, pd.DataFrame([{"hạng": "chung kết", "exp_id": r["exp_id"], "loại": "chung kết 3 seed",
                                           "mô tả": r["cấu hình"], "macro-F1 val": r["macro-F1 val"],
                                           "top-1 val": None, "macro-F1 test (mean ± std)": r["macro-F1 test"],
                                           "top-1 test (mean ± std)": r["top-1 test"]}])], ignore_index=True)
    return df


def write_xlsx(sheets: dict, path: Path, best_col: dict) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True)
            for j, col in enumerate(df.columns, start=1):
                letter = get_column_letter(j)
                width = max(10, min(60, int(max([len(str(col))] + [len(str(v)) for v in df[col].head(50)]) * 1.05)))
                ws.column_dimensions[letter].width = width
                if pd.api.types.is_float_dtype(df[col]):
                    fmt = "0.0" if ("ms" in col or "(s)" in col or "ảnh/s" in col or "s/epoch" in col) else "0.0000"
                    for cell in ws[letter][1:]:
                        cell.number_format = fmt
            col = best_col.get(name)
            if col and col in df.columns:
                vals = pd.to_numeric(df[col], errors="coerce")
                if vals.notna().any():
                    i = int(vals.idxmax()) + 2
                    for cell in ws[i]:
                        cell.fill = PatternFill("solid", fgColor="C6EFCE")


def backbone_plots(bk: pd.DataFrame, figs: Path) -> None:
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.5))
    for ax, xcol, xl in ((axes[0], "độ trễ batch-1 p50 (ms)", "độ trễ p50 batch 1 (ms)"),
                         (axes[1], "#tham số (M)", "số tham số (triệu)")):
        ax.scatter(bk[xcol], bk["macro-F1 val"])
        for _, r in bk.iterrows():
            ax.annotate(f"{r['exp_id']} {r['backbone'].split('_')[0]}", (r[xcol], r["macro-F1 val"]), fontsize=8,
                        xytext=(3, 3), textcoords="offset points")
        ax.set_xlabel(xl)
        ax.set_ylabel("macro-F1 val")
        ax.grid(alpha=0.3)
    fig.suptitle("Bước 1: macro-F1 val theo độ trễ và số tham số (1 seed, công thức T00)")
    fig.tight_layout()
    fig.savefig(figs / "backbones_f1_latency_params.png", dpi=120)
    plt.close(fig)


def _md(df: pd.DataFrame, floatfmt: str = "{:.4f}") -> str:
    def f(v):
        if isinstance(v, float):
            return "" if np.isnan(v) else floatfmt.format(v)
        return str(v)
    head = "| " + " | ".join(map(str, df.columns)) + " |"
    sep = "|" + "---|" * len(df.columns)
    body = ["| " + " | ".join(f(v) for v in row) + " |" for row in df.itertuples(index=False)]
    return "\n".join([head, sep, *body])


def main(out: str) -> None:
    root = Path(out)
    dec = json.loads((root / "decisions.json").read_text())
    figs = root / "figures"
    figs.mkdir(parents=True, exist_ok=True)
    bk, tr, inf = backbones_sheet(root), training_sheet(root), inference_sheet(root)
    fin, pc, lat = final_sheet(root, dec), perclass_sheet(root), latency_sheet(root)
    summ = summary_sheet(bk, tr, inf, fin)
    sheets = {"Backbones": bk, "Training": tr, "Inference": inf, "Final": fin, "PerClass": pc,
              "Latency": lat, "Summary": summ}
    if (root / "bonus_shift.csv").exists():
        sheets["Bonus_Shift"] = pd.read_csv(root / "bonus_shift.csv")
    write_xlsx(sheets, root / "results.xlsx",
               {"Backbones": "macro-F1 val", "Training": "macro-F1 val", "Inference": "macro-F1 val",
                "Summary": "macro-F1 val"})
    backbone_plots(bk, figs)
    md = [f"## {k}\n\n{_md(v)}\n" for k, v in sheets.items()]
    md.append("## decisions.json\n\n```json\n" + json.dumps(dec, indent=2, ensure_ascii=False) + "\n```\n")
    (root / "report_tables.md").write_text("\n".join(md), encoding="utf-8")
    print("đã ghi", root / "results.xlsx")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    main(ap.parse_args().out)
